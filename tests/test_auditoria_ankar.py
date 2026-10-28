"""Regressoes da etapa 10A.1.1: leitor XLSX isolado do PostgreSQL."""
from __future__ import annotations

from datetime import datetime
import hashlib
import json

from openpyxl import Workbook

from app.services.auditoria_ankar_service import auditar_planilhas


def _salvar(path, preparar):
    wb = Workbook()
    ws = wb.active
    preparar(ws)
    wb.save(path)
    wb.close()
    return path


def _rota(ws):
    ws.title = "ROTA 1"
    ws["A1"] = "ROTA 1"
    ws["C3"] = "Oceano"
    ws["C4"] = "Profundezas"
    ws["E3"] = "TIPOS DE TERRENO"
    ws["E4"] = "Água"
    ws["I3"] = "TEMPO DE VIAGEM"
    ws["I4"] = "3 Dias"
    ws["L3"] = "ACESSOS"
    ws["L4"] = "Yrezon"
    ws["A6"], ws["B6"], ws["C6"] = "% Captura", "ÁGUA", "d100"
    for n, (pc, pokemon, intervalo) in enumerate([
        (75, "Tentacool", "1-55"),
        (25, "Pyukumuku", "56-60"),
        (90, "Skrelp", "60-75"),
        (40, "Mantyke", "76-100"),
    ], start=7):
        ws[f"A{n}"] = pc
        ws[f"B{n}"] = pokemon
        ws[f"C{n}"] = intervalo


def _cidade(ws):
    ws.title = "GYM 01 - PSY"
    ws["A1"] = "CIDADE GINÁSIO 01 - YREZON"
    ws["C3"] = "Urbano"
    ws["C4"] = "Oceano"
    ws["E3"] = "TIPOS DE TERRENO"
    ws["E4"] = "Cimento"
    ws["J3"] = "TAMANHO DA CIDADE"
    ws["J4"] = "Grande - 2 Zonas"
    ws["M3"] = "ACESSOS"
    ws["M4"] = "Rota 1"
    ws["A18"] = "ZONA 01 - CENTRO"
    ws["A20"], ws["B20"], ws["C20"] = "% Captura", "CIDADE", "d100"
    ws["A21"], ws["B21"], ws["C21"] = 80, "Abra", datetime(2020, 5, 1)
    ws["B22"], ws["C22"] = "x", "6-100"
    ws["A33"] = "ZONA 02 - PORTO"
    ws["A35"], ws["B35"], ws["C35"] = "% Captura", "CIDADE", "d100"
    ws["E35"], ws["F35"], ws["G35"] = "% Captura", "ÁGUA", "d100"
    ws["A36"], ws["B36"], ws["C36"] = 90, "Trubbish", "1-100"
    ws["E36"], ws["F36"], ws["G36"] = 75, "Tentacool", "1-100"


def _arquivos(tmp_path):
    rotas = _salvar(tmp_path / "rotas.xlsx", _rota)
    cidades = _salvar(tmp_path / "cidades.xlsx", _cidade)
    return rotas, cidades


def test_leitura_identifica_local_terreno_zona_e_metadados(tmp_path):
    rel = auditar_planilhas(*_arquivos(tmp_path))
    assert len(rel.locais) == 2
    assert len(rel.tabelas) == 4
    assert rel.locais[0].habitat_principal == "Oceano"
    assert rel.locais[0].tempo_ou_tamanho == "3 Dias"
    assert rel.locais[0].acessos == ["Yrezon"]
    assert [(t.terreno, t.zona) for t in rel.tabelas if t.aba.startswith("GYM")] == [
        ("CIDADE", "ZONA 01 - CENTRO"),
        ("CIDADE", "ZONA 02 - PORTO"),
        ("ÁGUA", "ZONA 02 - PORTO"),
    ]


def test_sobreposicao_d100_mostra_as_duas_especies_e_celulas(tmp_path):
    rel = auditar_planilhas(*_arquivos(tmp_path))
    a = next(a for a in rel.achados if a.codigo == "sobreposicao")
    assert a.celula == "C9"
    assert "Pyukumuku" in a.mensagem and "Skrelp" in a.mensagem
    assert "60" in a.mensagem
    assert a.nivel == "ERRO"


def test_data_excel_nao_e_convertida_e_preserva_valor_original(tmp_path):
    rel = auditar_planilhas(*_arquivos(tmp_path))
    achado = next(a for a in rel.achados if a.codigo == "data_excel")
    assert achado.aba == "GYM 01 - PSY" and achado.celula == "C21"
    assert "2020-05-01" in achado.valor_original
    entrada = next(t for t in rel.tabelas if t.cabecalho == "C20").entradas[0]
    assert entrada.inicio is None and entrada.fim is None


def test_x_e_aviso_nao_e_automaticamente_encontro_vazio(tmp_path):
    rel = auditar_planilhas(*_arquivos(tmp_path))
    x = [a for a in rel.achados if a.codigo == "marcador_x"]
    assert len(x) == 1 and x[0].nivel == "AVISO"
    assert rel.tabelas[1].entradas[1].pokemon == "x"


def test_duas_zonas_na_mesma_linha_sao_diferenciadas_pela_coluna(tmp_path):
    def preparar(ws):
        ws.title = "P 01 - CAVERNA"
        ws["A1"] = "CAVERNA MAGNÉTICA"
        ws["A16"], ws["E16"] = "ZONA 01", "ZONA 02"
        ws["A18"], ws["B18"], ws["C18"] = "% Captura", "CAVERNA", "d100"
        ws["E18"], ws["F18"], ws["G18"] = "% Captura", "CAVERNA", "d100"
        ws["A19"], ws["B19"], ws["C19"] = 50, "Gastly", "1-100"
        ws["E19"], ws["F19"], ws["G19"] = 50, "Zubat", "1-100"
    rotas, cidades = _arquivos(tmp_path)
    _salvar(cidades, preparar)
    rel = auditar_planilhas(rotas, cidades)
    assert [(t.zona, t.terreno) for t in rel.tabelas if t.aba.startswith("P 01")] == [
        ("ZONA 01", "CAVERNA"), ("ZONA 02", "CAVERNA")
    ]


def test_intervalos_invalidos_lacunas_e_campos_incompletos(tmp_path):
    def preparar(ws):
        ws.title = "ROTA 3"
        ws["A1"] = "ROTA 3"
        ws["A6"], ws["B6"], ws["C6"] = "% Captura", "GRAMA", "d100"
        ws["A7"], ws["B7"], ws["C7"] = 75, "Pidgey", "0-5"
        ws["A8"], ws["B8"], ws["C8"] = 90, "Rattata", "30-10"
        ws["A9"], ws["B9"], ws["C9"] = 80, "Hoothoot", "texto ruim"
        ws["A10"], ws["B10"] = 90, "Sentret"
        ws["C11"] = "96-100"
    rotas, cidades = _arquivos(tmp_path)
    _salvar(rotas, preparar)
    rel = auditar_planilhas(rotas, cidades)
    codigos = {a.codigo for a in rel.achados}
    assert {"fora_d100", "faixa_invertida", "d100_invalido", "d100_ausente",
            "pokemon_ausente", "lacunas_d100"} <= codigos


def test_aba_sem_tabelas_e_relatorio_serializavel(tmp_path):
    rotas, cidades = _arquivos(tmp_path)
    _salvar(cidades, lambda ws: setattr(ws, "title", "P 07 - RESERVA CENTRAL"))
    rel = auditar_planilhas(rotas, cidades)
    assert any(a.codigo == "sem_tabelas" and a.nivel == "AVISO" for a in rel.achados)
    doc = json.loads(json.dumps(rel.como_dict()))
    assert doc["modo"] == "somente_leitura"
    assert "Resultado" in rel.texto(limite=None)


def test_auditoria_nao_modifica_arquivos_origem(tmp_path):
    rotas, cidades = _arquivos(tmp_path)
    antes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (rotas, cidades)]
    auditar_planilhas(rotas, cidades)
    depois = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (rotas, cidades)]
    assert depois == antes


def test_comando_flask_funciona_sem_banco_de_dados(tmp_path):
    import pytest
    pytest.importorskip("flask")
    from flask import Flask
    from app.commands.ankar import ankar_cli
    rotas, cidades = _arquivos(tmp_path)
    app = Flask(__name__)
    app.register_blueprint(ankar_cli, cli_group=None)
    runner = app.test_cli_runner()
    cmd = ["ankar", "auditar-planilhas", "--rotas", str(rotas), "--cidades", str(cidades)]
    resposta = runner.invoke(args=cmd)
    assert resposta.exit_code == 1  # overlap e data invalida
    assert "ERRO(S)" in resposta.output
    resposta_json = runner.invoke(args=cmd + ["--formato", "json"])
    assert resposta_json.exit_code == 1
    assert json.loads(resposta_json.output)["resumo"]["erros"] >= 2


def test_atualizacao_documentacao_preserva_historico_e_e_idempotente(tmp_path):
    import importlib.util
    from pathlib import Path

    base = Path(__file__).resolve().parents[1]
    caminho = base / "scripts" / "aplicar_documentacao_10a1.py"
    spec = importlib.util.spec_from_file_location("doc_ankar_10a1", caminho)
    modulo = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(modulo)
    docs = tmp_path / "docs"
    docs.mkdir()
    for nome in ("CONTINUIDADE_PROJETO.md", "DIARIO_DESENVOLVIMENTO.md"):
        (docs / nome).write_text("# Historico original\n\nNao apagar.\n", encoding="utf-8")
    modulo.aplicar(tmp_path, simular=True)
    assert all("10A.1.1" not in p.read_text(encoding="utf-8") for p in docs.iterdir())
    modulo.aplicar(tmp_path)
    conteudos = [p.read_text(encoding="utf-8") for p in docs.iterdir()]
    assert all(s.startswith("# Historico original\n\nNao apagar.") for s in conteudos)
    assert all(s.count(modulo.MARCADOR) == 1 for s in conteudos)
    modulo.aplicar(tmp_path)
    assert [p.read_text(encoding="utf-8") for p in docs.iterdir()] == conteudos


def test_formula_d100_e_preservada_para_revisao(tmp_path):
    rotas, cidades = _arquivos(tmp_path)
    from openpyxl import load_workbook
    w = load_workbook(rotas)
    w.active["C7"] = "=1+2"
    w.save(rotas)
    w.close()
    rel = auditar_planilhas(rotas, cidades)
    achado = next(a for a in rel.achados if a.codigo == "formula_d100")
    assert achado.celula == "C7"
    assert achado.valor_original == "'=1+2'"
